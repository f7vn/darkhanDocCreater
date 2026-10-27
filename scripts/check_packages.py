"""Generate and validate all five packages locally or through the deployed API."""
from __future__ import annotations

import argparse
from decimal import Decimal
import hashlib
from http.cookiejar import CookieJar
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
import pymupdf as fitz

import app


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base-url', help='Omit to call the local generator directly')
    parser.add_argument('--large-prices', action='store_true')
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()
    client = build_opener(HTTPCookieProcessor(CookieJar()))
    base = args.base_url.rstrip('/') if args.base_url else None
    if base:
        credentials = {key: os.environ.get(f'DOCGEN_{key.upper()}', '') for key in ('username', 'password')}
        if not all(credentials.values()):
            raise ValueError('Set DOCGEN_USERNAME and DOCGEN_PASSWORD')
        with client.open(base + '/health', timeout=30) as response:
            assert response.status == 200
        with client.open(Request(base + '/login', data=urlencode(credentials).encode()), timeout=30) as response:
            assert response.geturl().rstrip('/') == base

    area = Decimal('12.34') if args.large_prices else Decimal('32.5')
    area_price = Decimal('12345678.90') if args.large_prices else Decimal('21800.75')
    count = 2 if args.large_prices else 3
    stand_price = Decimal('87654321') if args.large_prices else Decimal('90000.50')
    markups = (Decimal('8.5'), Decimal('17.25')) if args.large_prices else (Decimal('7'), Decimal('15'))
    report = {}
    for mode in app.KTRU:
        payload = {'ktru': mode, 'area': str(area), 'areaPrice': str(area_price),
                   'standCount': str(count), 'standPrice': str(stand_price),
                   'markup1': str(markups[0]), 'markup2': str(markups[1])}
        if base:
            request = Request(base + '/api/generate', data=json.dumps(payload).encode(), headers={'Content-Type': 'application/json'})
            with client.open(request, timeout=180) as response:
                assert response.status == 201
                result = json.load(response)
        else:
            result = app.generate_from_payload(payload)
        assert len(result['files']) == 2
        files = {}
        for entry in result['files']:
            if base:
                with client.open(base + entry['url'], timeout=60) as response:
                    assert response.status == 200
                    assert response.headers['Content-Disposition'].startswith('attachment;')
                    files[entry['kind']] = response.read()
            else:
                files[entry['kind']] = (app.GENERATED_DIR / result['jobId'] / entry['name']).read_bytes()
            if args.output_dir:
                destination = args.output_dir / mode
                destination.mkdir(parents=True, exist_ok=True)
                (destination / entry['name']).write_bytes(files[entry['kind']])
        specification = Document(BytesIO(files['DOCX']))
        assert not specification._element.xpath('.//w:highlight'), mode
        assert len(specification.tables) == 2, mode
        for index, table in enumerate(specification.tables):
            if mode == 'combined':
                assert table.cell(1, 3).text == app.decimal_label(area)
                assert table.cell(2, 3).text == str(count)
            else:
                expected = f'{count} {"штук" if index == 0 else "дана"}' if mode == 'stands' else f'{app.decimal_label(area)} м2'
                assert table.cell(1, 2).text == expected, (mode, expected)
        pdf = fitz.open(stream=files['PDF'], filetype='pdf')
        assert 4 <= len(pdf) <= 15, (mode, len(pdf))
        text = re.sub(r'\s', '', '\n'.join(page.get_text() for page in pdf))
        keys = app.supplier_keys(mode)
        rates = app.supplier_rates(area_price, markups, keys)
        stands = app.supplier_rates(stand_price, markups, keys)
        for key in keys:
            expected = app.money(area * rates[key]) if mode != 'stands' else Decimal(0)
            if mode in ('stands', 'combined'):
                expected += app.money(count * stands[key])
            assert app.decimal_label(expected).replace(' ', '') in text, (mode, key, expected)
        # Check the actual offer pages, not only the totals in the justification.
        # All source image pixels, including handwritten ink and stamps, must
        # survive merging. Testing rendered text alone cannot catch lost art.
        signed = []
        for supplier in app.KTRU[mode]['suppliers']:
            if not supplier.get('signed_offer'):
                continue
            with fitz.open(app.TEMPLATE_DIR / supplier['signed_offer']) as original:
                expected_images = {
                    hashlib.sha256(fitz.Pixmap(original, item[0]).samples).hexdigest()
                    for item in original[0].get_images()
                }
            candidates = []
            for page in pdf:
                actual_images = {
                    hashlib.sha256(fitz.Pixmap(pdf, item[0]).samples).hexdigest()
                    for item in page.get_images()
                }
                if expected_images <= actual_images:
                    candidates.append(page)
            assert len(candidates) == 1, (mode, supplier['key'], 'signature/stamp image mismatch')
            page = candidates[0]
            offer_text = re.sub(r'\s', '', page.get_text())
            key = supplier['key']
            expected = app.money(area * rates[key]) if mode != 'stands' else Decimal(0)
            if mode in ('stands', 'combined'):
                expected += app.money(count * stands[key])
            assert app.decimal_label(expected).replace(' ', '') in offer_text, (mode, key, 'offer total')
            assert not list(page.widgets()), (mode, key, 'unexpected form field')
            assert 'SIGN' not in offer_text, (mode, key, 'signature marker')
            signed.append(key)
        report[mode] = {'jobId': result['jobId'], 'pdfPages': len(pdf), 'pdfBytes': len(files['PDF']), 'signedOffers': signed}
        print(mode, 'OK', report[mode], flush=True)
    if args.output_dir:
        (args.output_dir / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
