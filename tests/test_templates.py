import json
from decimal import Decimal
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

import app


EXPECTED = {
    'aesthetic': ('best', 'nalibai', 'archi'),
    'finishing': ('abdykhalykov', 'enzi', 'sarsenova'),
    'stands': ('abdicul', 'trade', 'andrey'),
    'graphics': ('best', 'nalibai', 'archi'),
    'combined': ('best', 'nalibai', 'archi'),
}


def text_of(document):
    return '\n'.join(p.text for p in app.iter_all_paragraphs(document))


def media_of(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.startswith('word/media/')}


class TemplateTests(unittest.TestCase):
    def test_five_separate_complete_source_sets(self):
        self.assertEqual(set(app.KTRU), set(EXPECTED))
        for mode, keys in EXPECTED.items():
            with self.subTest(mode=mode):
                config = app.KTRU[mode]
                self.assertEqual(app.supplier_keys(mode), keys)
                paths = [config['specification'], config['justification'], *(s['offer'] for s in config['suppliers'])]
                self.assertEqual(len(set(paths)), 5)
                self.assertEqual(len(config['source_files']), 5)
                for relative in paths:
                    self.assertEqual(Path(relative).parts[0], mode)
                    self.assertTrue((app.TEMPLATE_DIR / relative).is_file())

    def test_each_payload_uses_its_own_supplier_prices(self):
        for mode, keys in EXPECTED.items():
            with self.subTest(mode=mode), patch.object(app, 'generate_package') as generate:
                app.generate_from_payload({
                    'ktru': mode, 'area': '32,5', 'areaPrice': '21800,75',
                    'standCount': '3', 'standPrice': '90000,50',
                })
                actual_mode, area, count, stands, walls = generate.call_args.args
                self.assertEqual(actual_mode, mode)
                if mode != 'stands':
                    self.assertEqual(tuple(walls), keys)
                    self.assertEqual(walls[keys[1]], Decimal('23326.80'))
                    self.assertEqual(walls[keys[2]], Decimal('25070.86'))
                if mode in ('stands', 'combined'):
                    self.assertEqual(tuple(stands), keys)
                    self.assertEqual(stands[keys[1]], Decimal('96300.54'))

    def test_ui_supplier_data_is_rendered_from_server_catalog(self):
        html = app.render_index()
        embedded = re.search(r'<script id="work-types" type="application/json">(.*?)</script>', html, re.S).group(1)
        self.assertEqual(json.loads(embedded), app.public_catalog())
        self.assertNotIn('__WORK_TYPES__', html)
        self.assertEqual(json.loads(embedded)['finishing']['suppliers'], ['ИП Абдыхалыков', 'ENZI', 'ИП Сарсенова'])
        self.assertNotIn('source_folder', embedded)

    def test_area_specifications_keep_their_complete_original_bodies(self):
        with TemporaryDirectory() as directory:
            for mode in ('aesthetic', 'finishing', 'graphics'):
                with self.subTest(mode=mode):
                    source = app.TEMPLATE_DIR / app.KTRU[mode]['specification']
                    target = Path(directory) / f'{mode}.docx'
                    app.update_specification(source, target, Decimal('32.5'), app.KTRU[mode])
                    original = Document(source)
                    actual = Document(target)
                    self.assertEqual([p.text for p in actual.paragraphs if p.text.strip()],
                                     [p.text for p in original.paragraphs if p.text.strip()])
                    self.assertEqual(len(actual.tables), 2)
                    self.assertEqual([t.cell(1, 2).text for t in actual.tables], ['32,5 м2'] * 2)
                    self.assertEqual(media_of(source), media_of(target))
                    self.assertFalse(actual._element.xpath('.//w:highlight'))

    def test_area_offers_fill_prices_and_preserve_signature_media(self):
        with TemporaryDirectory() as directory:
            for mode in ('aesthetic', 'finishing', 'graphics'):
                for supplier in app.KTRU[mode]['suppliers']:
                    source = app.TEMPLATE_DIR / supplier['offer']
                    if source.suffix != '.docx':
                        continue
                    with self.subTest(mode=mode, provider=supplier['key']):
                        target = Path(directory) / f'{mode}-{supplier["key"]}.docx'
                        app.update_area_offer(source, target, supplier['key'], Decimal('32.5'), Decimal('21800.75'))
                        document = Document(target)
                        text = text_of(document)
                        self.assertIn('32,5', text)
                        self.assertIn('21 800,75', text)
                        self.assertIn('708 524,38', text)
                        self.assertFalse(document._element.xpath('.//w:highlight'))
                        self.assertEqual(media_of(source), media_of(target))
                        self.assertFalse(document._element.xpath(".//wp:anchor[@behindDoc='1'][wp:extent[@cy > 10000000]]"))
                        if supplier['key'] == 'sarsenova':
                            for cell in document.tables[0].rows[0].cells:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        self.assertEqual(run.font.name, 'Arial')
                                        self.assertEqual(run.font.size.pt, 11)

    def test_all_area_justifications_have_current_bilingual_price_lists(self):
        with TemporaryDirectory() as directory:
            for mode in ('aesthetic', 'finishing', 'graphics', 'combined'):
                with self.subTest(mode=mode):
                    keys = app.supplier_keys(mode)
                    rates = app.supplier_rates(Decimal('21800.75'), (Decimal('7'), Decimal('15')), keys)
                    stands = app.supplier_rates(Decimal('90000.5'), (Decimal('7'), Decimal('15')), keys) if mode == 'combined' else None
                    target = Path(directory) / f'{mode}.docx'
                    app.update_catalog_justification(mode, target, Decimal('32.5'), rates, 3 if stands else 0, stands)
                    document = Document(target)
                    text = text_of(document)
                    for key in keys:
                        total = app.money(Decimal('32.5') * rates[key]) + (app.money(3 * stands[key]) if stands else 0)
                        self.assertGreaterEqual(text.count(app.decimal_label(total)), 2)
                        self.assertIn(app.company_name(key), text)
                        self.assertIn(app.company_name(key, True), text)
                    for stale in ('2\u00a0000 000', '2\u00a0120 000', '2\u00a0320 000', '990 000'):
                        self.assertNotIn(stale, text)
                    self.assertFalse(document._element.xpath('.//w:highlight'))
                    if mode == 'graphics':
                        self.assertNotIn('эстетическому оформлению составляет', text)
                    self.assertIn('что не превышает', text)
                    self.assertIn('аспайды', text)

    def test_large_totals_update_both_conclusions_without_editing_the_law_quote(self):
        with TemporaryDirectory() as directory:
            rates = {key: Decimal('30000') for key in app.supplier_keys('finishing')}
            target = Path(directory) / 'large.docx'
            app.update_catalog_justification('finishing', target, Decimal('100'), rates)
            document = Document(target)
            text = text_of(document)
            self.assertIn('что превышает', text)
            self.assertIn('асады', text)
            self.assertIn('3 000 000', text)
            original = Document(app.TEMPLATE_DIR / app.KTRU['finishing']['justification'])
            for index in (0, 1, 2, 3):
                self.assertEqual(document.tables[0].rows[index].cells[0].text, original.tables[0].rows[index].cells[0].text)


if __name__ == '__main__':
    unittest.main()
