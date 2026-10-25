import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import app
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt


class PricingTests(unittest.TestCase):
    def test_replaced_text_has_no_highlight_and_keeps_other_formatting(self):
        paragraph = Document().add_paragraph()
        original = paragraph.add_run('_____')
        original.font.highlight_color = WD_COLOR_INDEX.YELLOW
        original.font.name = 'Arial'
        original.font.size = Pt(14)
        original.bold = True

        app.set_paragraph_text(paragraph, '32')

        self.assertEqual(paragraph.text, '32')
        self.assertIsNone(paragraph.runs[0].font.highlight_color)
        self.assertEqual(paragraph.runs[0].font.name, 'Arial')
        self.assertEqual(paragraph.runs[0].font.size, Pt(14))
        self.assertTrue(paragraph.runs[0].bold)

    def test_combined_specification_has_no_highlight_in_either_language(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / 'specification.docx'
            app.update_combined_specification(
                app.TEMPLATE_DIR / app.KTRU['combined']['specification'], destination,
                Decimal('32'), 1,
            )
            document = Document(destination)
            self.assertEqual(len(document.tables), 2)
            for table in document.tables:
                self.assertEqual(table.cell(1, 3).text, '32')
                self.assertEqual(table.cell(2, 3).text, '1')
            self.assertFalse(document._element.xpath('.//w:highlight'))

    def test_combined_justification_has_no_highlight(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / 'justification.docx'
            app.update_combined_justification(
                app.TEMPLATE_DIR / app.KTRU['combined']['justification'], destination,
                Decimal('900000'),
            )
            document = Document(destination)
            self.assertFalse(document._element.xpath('.//w:highlight'))
            text = '\n'.join(p.text for p in app.iter_all_paragraphs(document))
            self.assertEqual(text.count('900 000'), 2)

    def test_stands_justification_has_no_highlight(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / 'justification.docx'
            app.update_stands_justification(destination, 3, {
                'abdicul': Decimal('100000'),
                'trade': Decimal('107000'),
                'andrey': Decimal('115000'),
            })
            self.assertFalse(Document(destination)._element.xpath('.//w:highlight'))

    def test_nalibai_background_removed_and_stamp_preserved(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / 'offer.docx'
            app.update_combined_offer_nalibai(
                app.TEMPLATE_DIR / 'combined/offer-nalibai.docx', destination,
                Decimal('80'), 2, Decimal('107000'), Decimal('2354000'), Decimal('26750'),
            )
            document = Document(destination)
            self.assertFalse(document._element.xpath(".//wp:anchor[@behindDoc='1']"))
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertIn('2 354 000', document.tables[0].cell(4, 1).text)

    def test_localized_number_grouping(self):
        self.assertEqual(app.parse_area('1\u00a0234,5'), Decimal('1234.5'))
        self.assertEqual(app.parse_count('1\u202f000'), 1000)
        self.assertEqual(app.parse_money('100\u00a0000,50', 'цену'), Decimal('100000.50'))
        self.assertEqual(app.parse_markup('1\u00a0000'), Decimal('1000'))

    def test_default_competitor_percentages(self):
        actual = app.supplier_rates(Decimal('100000'), (Decimal('7'), Decimal('15')), ('base', 'first', 'second'))
        self.assertEqual(actual, {'base': Decimal('100000'), 'first': Decimal('107000'), 'second': Decimal('115000')})

    def test_round_half_up_to_cents(self):
        actual = app.supplier_rates(Decimal('0.10'), (Decimal('5'), Decimal('15')), ('base', 'first', 'second'))
        self.assertEqual(actual['first'], Decimal('0.11'))
        self.assertEqual(actual['second'], Decimal('0.12'))

    def test_stands_need_no_area_and_use_stand_suppliers(self):
        with patch.object(app, 'generate_package') as generate:
            app.generate_from_payload({'ktru': 'stands', 'standCount': '3', 'standPrice': '100 000'})
        mode, area, count, rates, _ = generate.call_args.args
        self.assertEqual((mode, area, count), ('stands', Decimal(0), 3))
        self.assertEqual(rates, {'abdicul': Decimal('100000'), 'trade': Decimal('107000'), 'andrey': Decimal('115000')})

    def test_custom_markup_applies_to_both_components(self):
        with patch.object(app, 'generate_package') as generate:
            app.generate_from_payload({'ktru': 'combined', 'area': '12,5', 'areaPrice': '20000', 'standCount': '2', 'standPrice': '100000', 'markup1': '8,5', 'markup2': '20'})
        _, area, count, stands, walls = generate.call_args.args
        self.assertEqual((area, count), (Decimal('12.5'), 2))
        self.assertEqual(stands['nalibai'], Decimal('108500'))
        self.assertEqual(walls['nalibai'], Decimal('21700'))
        self.assertEqual(stands['archi'], Decimal('120000'))
        self.assertEqual(walls['archi'], Decimal('24000'))

    def test_invalid_prices_and_percentages(self):
        for value in ('NaN', 'Infinity', '-1', '0', '0.001', '100000001'):
            with self.subTest(price=value), self.assertRaises(ValueError):
                app.parse_money(value, 'цену за стенд')
        for value in ('NaN', 'Infinity', '-1', '1001', '1.111'):
            with self.subTest(markup=value), self.assertRaises(ValueError):
                app.parse_markup(value)


if __name__ == '__main__':
    unittest.main()
