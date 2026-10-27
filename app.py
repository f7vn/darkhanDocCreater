from __future__ import annotations

import json
import hmac
import os
import re
import secrets
import shutil
import subprocess
import tempfile
import uuid
from copy import deepcopy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http import HTTPStatus
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

import pymupdf as fitz
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
TEMPLATE_DIR = ROOT / "templates"
GENERATED_DIR = ROOT / "generated"
GENERATED_DIR.mkdir(exist_ok=True)

AUTH_USERNAME = os.environ.get("DOCGEN_USERNAME", "")
AUTH_PASSWORD = os.environ.get("DOCGEN_PASSWORD", "")
AUTH_COOKIE_NAME = "docgen_session"
AUTH_SESSION_TOKEN = secrets.token_urlsafe(32)
AUTH_SESSION_SECONDS = 8 * 60 * 60

KTRU = json.loads((TEMPLATE_DIR / "catalog.json").read_text(encoding="utf-8"))


def supplier_keys(ktru_key: str) -> tuple[str, str, str]:
    return tuple(supplier["key"] for supplier in KTRU[ktru_key]["suppliers"])


def public_catalog() -> dict:
    return {
        key: {"label": kind["label"], "suppliers": [s["name"] for s in kind["suppliers"]]}
        for key, kind in KTRU.items()
    }


def render_index() -> str:
    catalog_json = json.dumps(public_catalog(), ensure_ascii=False).replace("<", "\\u003c")
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8").replace("__WORK_TYPES__", catalog_json)

CAP_TOTAL = Decimal("2162500")


def decimal_label(value: Decimal) -> str:
    if value == value.to_integral():
        return f"{int(value):,}".replace(",", " ")
    normalized = format(value.normalize(), "f")
    integer, fraction = normalized.split(".")
    return f"{int(integer):,}".replace(",", " ") + "," + fraction


def slug_area(value: Decimal) -> str:
    return decimal_label(value).replace(" ", "").replace(",", "-")


def parse_area(raw: object) -> Decimal:
    text = re.sub(r"\s", "", str(raw)).replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Введите площадь числом") from exc
    if not value.is_finite() or value <= 0:
        raise ValueError("Площадь должна быть больше нуля")
    if value > Decimal("100000"):
        raise ValueError("Площадь не может превышать 100 000 м²")
    if value.as_tuple().exponent < -2:
        raise ValueError("Используйте не более двух знаков после запятой")
    return value


def parse_count(raw: object) -> int:
    text = re.sub(r"\s", "", str(raw))
    if not re.fullmatch(r"\d+", text):
        raise ValueError("Введите количество стендов целым числом")
    value = int(text)
    if value <= 0:
        raise ValueError("Количество стендов должно быть больше нуля")
    if value > 10000:
        raise ValueError("Количество стендов не может превышать 10 000")
    return value


def parse_money(raw: object, label: str) -> Decimal:
    text = re.sub(r"\s", "", str(raw)).replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Укажите {label}") from exc
    if not value.is_finite() or value <= 0:
        raise ValueError(f"{label.capitalize()} должна быть больше нуля")
    if value > Decimal("100000000"):
        raise ValueError(f"{label.capitalize()} слишком велика")
    if value.as_tuple().exponent < -2:
        raise ValueError("В цене можно использовать не более двух знаков после запятой")
    return value


def parse_markup(raw: object) -> Decimal:
    try:
        value = Decimal(re.sub(r"\s", "", str(raw)).replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError("Введите надбавку конкурента в процентах") from exc
    if not value.is_finite() or value < 0 or value > 1000:
        raise ValueError("Надбавка должна быть от 0 до 1 000 процентов")
    if value.as_tuple().exponent < -2:
        raise ValueError("В надбавке можно использовать не более двух знаков после запятой")
    return value


def money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def supplier_rates(base: Decimal, markups: tuple[Decimal, Decimal], keys: tuple[str, str, str]) -> dict[str, Decimal]:
    return {
        key: money(base * (Decimal(1) + percent / 100))
        for key, percent in zip(keys, (Decimal(0), *markups))
    }


def iter_all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def set_paragraph_text(paragraph, text: str, bold_prefix: str | None = None) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)
    # Placeholder highlighting belongs to the template, not the final values.
    run.font.highlight_color = None
    if run.font.name is None:
        run.font.name = "Times New Roman"
    if run.font.size is None:
        run.font.size = Pt(11)
    if bold_prefix and text.startswith(bold_prefix):
        # Keep the source layout stable; the provider remains easy to scan through numbering.
        run.bold = False


def clear_highlighting(document: Document) -> None:
    for node in document._element.xpath('.//w:highlight'):
        node.getparent().remove(node)


def remove_page_backgrounds(document: Document) -> None:
    """Remove only page-sized gray paper backgrounds, never logos or stamps."""
    section = document.sections[0]
    for anchor in document._element.xpath(".//wp:anchor[@behindDoc='1']"):
        extent = anchor.find(qn('wp:extent'))
        if extent is not None and (
            int(extent.get('cx', '0')) >= section.page_width * 0.9
            and int(extent.get('cy', '0')) >= section.page_height * 0.9
        ):
            drawing = anchor.getparent()
            drawing.getparent().remove(drawing)


def remove_empty_spacers(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        if (not paragraph.text.strip()
                and not paragraph._p.xpath('.//w:drawing | .//w:pict | .//w:sectPr')):
            paragraph._element.getparent().remove(paragraph._element)


def company_name(key: str, kazakh: bool = False) -> str:
    names = {
        'best': ('ИП «Best Production»', '«Best Production» ЖК'),
        'nalibai': ('ИП «Нәлібай»', '«Нәлібай» ЖК'),
        'archi': ('ТОО «Архиспецснабмонтаж»', '«Архиспецснабмонтаж» ЖШС'),
        'abdykhalykov': ('ИП «Абдыхалыков»', '«Абдыхалыков» ЖК'),
        'enzi': ('ТОО «ЕНЗИ» (ENZI)', '«ЕНЗИ» (ENZI) ЖШС'),
        'sarsenova': ('ИП «Сарсенова»', '«Сарсенова» ЖК'),
        'abdicul': ('ИП «Әбдіқұл»', '«Әбдіқұл» ЖК'),
        'trade': ('ТОО «Торговая сеть Казахстана»', '«Торговая сеть Казахстана» ЖШС'),
        'andrey': ('ИП «Андрей и Ко»', '«Андрей и Ко» ЖК'),
    }
    return names[key][int(kazakh)]


def update_catalog_justification(
    ktru_key: str, destination: Path, area: Decimal,
    area_rates: dict[str, Decimal], count: int = 0,
    stand_rates: dict[str, Decimal] | None = None,
) -> None:
    """Fill prices in the selected original, including table-cell paragraphs."""
    config = KTRU[ktru_key]
    doc = Document(TEMPLATE_DIR / config['justification'])
    keys = supplier_keys(ktru_key)
    totals = {key: money(area * area_rates[key]) + money(count * (stand_rates or {}).get(key, Decimal(0))) for key in keys}
    minimum_key = min(keys, key=totals.get)
    minimum = totals[minimum_key]
    anchors = []
    conclusions = 0
    for paragraph in list(iter_all_paragraphs(doc)):
        text = paragraph.text.strip()
        if text.startswith(('Для определения среднерыночной', 'Ұқсас қызметтердің/жұмыстардың орташа')):
            anchors.append((paragraph, text.startswith('Ұқсас')))
        if text.startswith(('По итогам проведённого анализа', 'Жүргізілген талдау нәтижесінде')):
            updated, changes = re.subn(
                r'(составляет\s+|алмағанда\s+)[\d\s_]+(?=тенге|теңгені)',
                lambda match: match.group(1) + decimal_label(minimum) + ' ', text,
            )
            if changes != 1:
                raise ValueError(f'Не найдено поле итоговой цены в шаблоне {ktru_key}')
            if minimum > CAP_TOTAL:
                updated = updated.replace('не превышает', 'превышает').replace('аспайды', 'асады')
            if ktru_key == 'graphics':
                # Folders 1 and 4 contain the same justification with an aesthetic
                # subject. Keep its source layout but correct this subject slot.
                updated = updated.replace('эстетическому оформлению', 'нанесению изображений и надписей на стены')
                updated = updated.replace('эстетикалық безендіру', 'қабырғаларға кескіндер мен жазулар түсіру')
            set_paragraph_text(paragraph, updated)
            conclusions += 1
        elif 'наименьшая' in text and 'определена' in text:
            _, suffix = text.split('что подтверждает', 1)
            set_paragraph_text(paragraph,
                f'В результате проведенного анализа цен наименьшая цена определена у '
                f'{company_name(minimum_key)} — {decimal_label(minimum)} тенге без НДС, что подтверждает{suffix}')
        elif 'ең төмен баға' in text and 'анықталды' in text:
            _, suffix = text.split('бұл бюджет', 1)
            set_paragraph_text(paragraph,
                f'Жүргізілген баға талдауының нәтижесінде ең төмен баға '
                f'{company_name(minimum_key, True)} бойынша {decimal_label(minimum)} теңге '
                f'мөлшерінде анықталды, бұл бюджет{suffix}')
    if conclusions != 2 or len(anchors) != 2:
        raise ValueError(f'Неполный двуязычный шаблон обоснования: {ktru_key}')

    for anchor, kazakh in anchors:
        # Delete the old illustrative provider prices, not the legal or market
        # reference text in the other rows. Missing price lists are added here.
        parent = anchor._parent
        for paragraph in list(parent.paragraphs):
            text = paragraph.text.strip()
            if paragraph._p is anchor._p:
                continue
            if not text or (('×' in text or '*' in text) and any(name in text for name in ('Production', 'Нәлібай', 'Архиспецснабмонтаж'))):
                paragraph._element.getparent().remove(paragraph._element)
        previous = anchor
        for key in keys:
            parts = [f'{decimal_label(area)} м² × {decimal_label(area_rates[key])}']
            if stand_rates is not None:
                parts.append(f'{count} {"дана" if kazakh else "шт."} × {decimal_label(stand_rates[key])}')
            text = f'{company_name(key, kazakh)}: ' + ' + '.join(parts) + f' = {decimal_label(totals[key])} {"теңге" if kazakh else "тенге"}.'
            element = OxmlElement('w:p')
            if anchor._p.pPr is not None:
                element.append(deepcopy(anchor._p.pPr))
            previous._p.addnext(element)
            previous = Paragraph(element, parent)
            set_paragraph_text(previous, text)
            previous.paragraph_format.space_before = Pt(0)
            previous.paragraph_format.space_after = Pt(4)
    clear_highlighting(doc)
    doc.save(destination)


def update_area_offer(source: Path, destination: Path, provider: str, area: Decimal, rate: Decimal) -> None:
    doc = Document(source)
    total = money(area * rate)
    remove_page_backgrounds(doc)
    if provider == 'best':
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith('Общая площадь'):
                set_paragraph_text(paragraph,
                    f'Общая площадь печати: {decimal_label(area)} м²\n'
                    f'Стоимость работ за 1 м²: {decimal_label(rate)} тенге\n'
                    f'Общая стоимость: {decimal_label(total)} тенге')
        # Keep the original letterhead and signature, removing only spacer lines.
        remove_empty_spacers(doc)
    elif provider == 'nalibai':
        table = doc.tables[0]
        set_paragraph_text(table.cell(1, 3).paragraphs[0], f'{decimal_label(area)} м²')
        set_paragraph_text(table.cell(1, 4).paragraphs[0], f'{decimal_label(total)} тенге')
        set_paragraph_text(table.cell(2, 4).paragraphs[0], f'{decimal_label(total)} тенге')
        paragraph = table.cell(1, 2).paragraphs[0]
        set_paragraph_text(paragraph, paragraph.text + f'\nЦена за 1 м²: {decimal_label(rate)} тенге')
        remove_empty_spacers(doc)
    elif provider == 'abdykhalykov':
        table = doc.tables[1]
        set_paragraph_text(table.cell(1, 1).paragraphs[0], f'Декоративно-отделочные работы\n{decimal_label(rate)} тенге за 1 м²')
        set_paragraph_text(table.cell(1, 2).paragraphs[0], f'{decimal_label(area)} м²')
        set_paragraph_text(table.cell(1, 3).paragraphs[0], f'{decimal_label(total)} тенге')
    elif provider in ('enzi', 'sarsenova'):
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text.startswith(('Площадь:', 'Общая площадь:')):
                set_paragraph_text(paragraph, f'Общая площадь: {decimal_label(area)} м².')
            elif text.startswith('Стоимость работ за'):
                set_paragraph_text(paragraph, f'Стоимость работ за 1 м²: {decimal_label(rate)} тенге.')
            elif text.startswith('Стоимость:'):
                set_paragraph_text(paragraph, f'Цена за 1 м²: {decimal_label(rate)} тенге.\nСтоимость: {decimal_label(total)} тенге без НДС.')
            elif text.startswith('Итоговая стоимость:'):
                set_paragraph_text(paragraph, f'Итоговая стоимость: {decimal_label(total)} тенге.')
        if provider == 'sarsenova':
            remove_empty_spacers(doc)
            # The source header inherits a theme font; Linux substitutes it
            # differently from Word/macOS. Explicit metrics keep the address
            # and bank details inside the original letterhead background.
            for cell in doc.tables[0].rows[0].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(11)
            for paragraph in doc.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
            # The decorative header has a fixed page anchor, independent of
            # the table height. Leave a gap after the now-compact requisites.
            doc.paragraphs[0].paragraph_format.space_before = Pt(18)
            for cell in doc.tables[-1].rows[0].cells:
                for paragraph in list(cell.paragraphs):
                    if not paragraph.text.strip() and not paragraph._p.xpath('.//w:drawing | .//w:pict'):
                        paragraph._element.getparent().remove(paragraph._element)
    else:
        raise ValueError(f'Нет обработчика шаблона КП: {provider}')
    clear_highlighting(doc)
    doc.save(destination)


def create_area_offer_archi(source_path: Path, destination: Path, area: Decimal, rate: Decimal) -> None:
    source = fitz.open(source_path)
    page = source[0]
    erase(page, fitz.Rect(400, 1890, 2310, 2040))
    textbox(page, fitz.Rect(419, 1905, 2290, 1965),
            f'Объем работ: {decimal_label(area)} м² × {decimal_label(rate)} тенге / м²', 48, 30)
    textbox(page, fitz.Rect(419, 1968, 2290, 2035),
            f'Общая стоимость работ: {decimal_label(money(area * rate))} тенге', 48, 30, bold=True)
    save_flattened_a4(source, destination)
    source.close()




def update_specification(source: Path, destination: Path, area: Decimal, kind: dict) -> None:
    doc = Document(source)
    # Each work type supplies its own complete specification. Only fill its slots;
    # do not turn an aesthetic specification into another work by renaming it.
    if kind.get('specification', '').startswith('graphics/'):
        # The source's Kazakh part inherits wider line spacing than Russian.
        # Match the Russian spacing so its appendix stays in the same section.
        for paragraph in doc.paragraphs:
            paragraph.paragraph_format.line_spacing = 1.0
    area_text = f"{decimal_label(area)} м2"
    for table in doc.tables:
        if len(table.rows) > 1 and len(table.columns) > 2:
            set_paragraph_text(table.cell(1, 2).paragraphs[0], area_text)

    # The source leaves eight empty lines before the Russian appendix. In Word,
    # this can strand the header row at the foot of one page and the data row on
    # the next. Keep four lines of visual separation and make this one table a
    # little denser, without changing the main body or the Kazakh section.
    russian_appendix_match = next(
        (
            (index, paragraph)
            for index, paragraph in enumerate(doc.paragraphs)
            if paragraph.text.strip() == "Приложение 1"
        ),
        None,
    )
    if russian_appendix_match is not None:
        appendix_index, russian_appendix = russian_appendix_match
        paragraphs = doc.paragraphs
        empty_before = []
        for paragraph in reversed(paragraphs[:appendix_index]):
            if paragraph.text.strip():
                break
            empty_before.append(paragraph)
        for paragraph in empty_before[:4]:
            paragraph._element.getparent().remove(paragraph._element)

        russian_appendix.paragraph_format.space_before = Pt(0)
        russian_appendix.paragraph_format.space_after = Pt(3)
        russian_appendix.paragraph_format.keep_with_next = True

    # Replace the long run of spacer paragraphs between the two language
    # sections with a deterministic page break. This keeps the Kazakh heading
    # at the top of a fresh page in both Word and LibreOffice.
    kazakh_heading_match = next(
        (
            (index, paragraph)
            for index, paragraph in enumerate(doc.paragraphs)
            if paragraph.text.strip() == "ТЕХНИКАЛЫҚ СИПАТТАМА"
        ),
        None,
    )
    if kazakh_heading_match is not None:
        heading_index, kazakh_heading = kazakh_heading_match
        paragraphs = doc.paragraphs
        empty_before = []
        for paragraph in reversed(paragraphs[:heading_index]):
            if paragraph.text.strip():
                break
            empty_before.append(paragraph)
        for paragraph in empty_before:
            paragraph._element.getparent().remove(paragraph._element)
        kazakh_heading.paragraph_format.page_break_before = True
        kazakh_heading.paragraph_format.space_before = Pt(0)

    kazakh_appendix_match = next(
        (
            (index, paragraph)
            for index, paragraph in enumerate(doc.paragraphs)
            if paragraph.text.strip() == "Қосымша 1"
        ),
        None,
    )
    if kazakh_appendix_match is not None:
        appendix_index, kazakh_appendix = kazakh_appendix_match
        paragraphs = doc.paragraphs
        empty_before = []
        for paragraph in reversed(paragraphs[:appendix_index]):
            if paragraph.text.strip():
                break
            empty_before.append(paragraph)
        for paragraph in empty_before:
            paragraph._element.getparent().remove(paragraph._element)
        kazakh_appendix.paragraph_format.space_before = Pt(0)
        kazakh_appendix.paragraph_format.space_after = Pt(3)
        kazakh_appendix.paragraph_format.keep_with_next = True

    for specification_table in doc.tables:
        # The reference stores both appendix tables as floating objects. Word
        # may then place the caption and table on different pages. Inline
        # tables paginate predictably and preserve the same visible geometry.
        floating_position = specification_table._tbl.tblPr.find(qn("w:tblpPr"))
        if floating_position is not None:
            specification_table._tbl.tblPr.remove(floating_position)

        for row_index, row in enumerate(specification_table.rows):
            row.height = Pt(28)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.keep_together = True
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs:
                        run.font.size = Pt(12)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() in ('Приложение 1', 'Қосымша 1', '1-қосымша'):
            paragraph.paragraph_format.keep_with_next = True
            following = paragraph._p.getnext()
            while following is not None and following.tag == qn('w:p') and not following.xpath('.//w:t | .//w:drawing | .//w:pict | .//w:sectPr'):
                next_element = following.getnext()
                following.getparent().remove(following)
                following = next_element
    doc.save(destination)


def arial_font(bold: bool = False) -> Path:
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise RuntimeError("Не найден шрифт с поддержкой кириллицы (Arial или DejaVu Sans)")


def erase(page, *rects: fitz.Rect) -> None:
    for rect in rects:
        page.add_redact_annot(rect, fill=(1, 1, 1), cross_out=False)
    page.apply_redactions(
        images=fitz.PDF_REDACT_IMAGE_NONE,
        graphics=fitz.PDF_REDACT_LINE_ART_NONE,
        text=fitz.PDF_REDACT_TEXT_REMOVE,
    )


def textbox(page, rect: fitz.Rect, text: str, max_size: float, min_size: float = 5.0,
            bold: bool = False, align: int = fitz.TEXT_ALIGN_LEFT) -> None:
    font_path = arial_font(bold)
    font_name = "DocArialBold" if bold else "DocArial"
    page.insert_font(fontname=font_name, fontfile=str(font_path))
    size = max_size
    while size >= min_size:
        shape = page.new_shape()
        result = shape.insert_textbox(
            rect,
            text,
            fontname=font_name,
            fontfile=str(font_path),
            fontsize=size,
            lineheight=1.05,
            color=(0, 0, 0),
            align=align,
        )
        if result >= 0:
            shape.commit(overlay=True)
            return
        size -= 0.4
    raise RuntimeError(f"Текст не помещается в шаблон: {text}")


def pdf_field(page, bounds: tuple, text: str, size: float = 11, minimum: float = 7,
              bold: bool = False, align: int = fitz.TEXT_ALIGN_LEFT,
              fill: tuple = (1, 1, 1)) -> None:
    """Replace only a verified variable slot, leaving images and table rules intact."""
    rect = fitz.Rect(bounds)
    # No redaction fill: MuPDF paints it with a 1pt outline which can cover
    # adjacent table rules. Paint the exact field rectangle ourselves below.
    page.add_redact_annot(rect, fill=None, cross_out=False)
    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                          graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                          text=fitz.PDF_REDACT_TEXT_REMOVE)
    # The source's yellow placeholder highlights are vector drawings, not
    # annotations. An opaque field fill removes them without touching stamps.
    page.draw_rect(rect, color=None, fill=fill, overlay=True)
    textbox(page, rect, text, size, minimum, bold=bold, align=align)


def create_signed_offer(source_path: Path, destination: Path, mode: str, provider: str,
                        area: Decimal, area_rate: Decimal, count: int,
                        stand_rate: Decimal) -> None:
    """Fill the nine approved PDF originals without reflowing their signed artwork."""
    area_total = money(area * area_rate)
    stand_total = money(count * stand_rate)
    total = area_total + stand_total
    with fitz.open(source_path) as source:
        if len(source) != 1:
            raise ValueError(f'Ожидался одностраничный шаблон КП: {source_path.name}')
        page = source[0]
        for widget in list(page.widgets()):
            # Preview adds an EMPTY /Sig widget displaying the orange SIGN
            # arrow. The actual handwritten signature is a page image. Never
            # strip a real cryptographic signature or a non-Apple form field.
            apple_marker = source.xref_get_key(widget.xref, 'AAPL:isSignatureMarker')[1]
            signature_value = source.xref_get_key(widget.xref, 'V')[0]
            if (widget.field_type != fitz.PDF_WIDGET_TYPE_SIGNATURE
                    or apple_marker != 'true' or signature_value != 'null'):
                raise ValueError('В шаблоне обнаружена электронная подпись или неизвестное поле PDF')
            page.delete_widget(widget)

        if provider in ('nalibai', 'abdykhalykov'):
            # Keep the background image (including the supplier logo), covering
            # just its gray 34pt left strip, well outside text and signature art.
            page.draw_rect(fitz.Rect(0, 0, 34, page.rect.height), color=None, fill=(1, 1, 1))

        if provider == 'best':
            top = 474.0 if mode == 'combined' else 444.3
            lines = [
                f'Общая площадь печати: {decimal_label(area)} м²',
                f'Стоимость работ за 1 м²: {decimal_label(area_rate)} тенге',
            ]
            if mode == 'combined':
                lines.append(f'Стенды: {count} шт. × {decimal_label(stand_rate)} = {decimal_label(stand_total)} тенге')
            lines.append(f'Общая стоимость: {decimal_label(total)} тенге')
            for index, line in enumerate(lines):
                y = top + 14.85 * index
                pdf_field(page, (46.5, y, 565, y + 14.7), line, 10.5, 8,
                          bold=index == len(lines) - 1)
        elif provider == 'nalibai' and mode != 'combined':
            bottom = 254.1 if mode == 'aesthetic' else 267.9
            pdf_field(page, (380, 194, 428, bottom), f'{decimal_label(area)}\nм²',
                      align=fitz.TEXT_ALIGN_CENTER)
            pdf_field(page, (432, 194, 527, bottom),
                      f'{decimal_label(area_rate)} тг/м²\nСумма:\n{decimal_label(total)} тенге',
                      10, 7, align=fitz.TEXT_ALIGN_CENTER)
            pdf_field(page, (432, bottom + 2, 527, bottom + 22.5),
                      f'{decimal_label(total)} тенге', 10, 6.5, bold=True,
                      align=fitz.TEXT_ALIGN_CENTER, fill=(0.851, 0.918, 0.969))
        elif provider == 'nalibai':
            pdf_field(page, (216, 420.6, 576, 434.1),
                      f'{decimal_label(area)} м² × {decimal_label(area_rate)} = {decimal_label(area_total)} тенге', 10, 7)
            pdf_field(page, (216, 435, 576, 462.2),
                      f'{count} шт. × {decimal_label(stand_rate)} = {decimal_label(stand_total)} тенге', 10, 7)
            pdf_field(page, (216, 463, 576, 490.3),
                      f'{decimal_label(total)} тенге без НДС', 11, 8, bold=True)
        elif provider == 'abdykhalykov':
            pdf_field(page, (330, 493.2, 426.5, 508.5), f'{decimal_label(area)} м²',
                      11, 8, align=fitz.TEXT_ALIGN_CENTER)
            pdf_field(page, (431, 493.2, 526, 508.5), f'{decimal_label(total)} тенге',
                      11, 6.5, align=fitz.TEXT_ALIGN_CENTER)
            pdf_field(page, (100, 512, 526, 530),
                      f'Стоимость работ за 1 м²: {decimal_label(area_rate)} тенге', 11, 9)
        elif provider == 'sarsenova':
            pdf_field(page, (56.5, 517.3, 559, 534.3),
                      f'Стоимость работ за 1 м²: {decimal_label(area_rate)} тенге.', 13, 10)
            pdf_field(page, (56.5, 534.5, 559, 551.5),
                      f'Общая площадь: {decimal_label(area)} м².', 13, 10)
            pdf_field(page, (56.5, 576.6, 559, 594),
                      f'Итоговая стоимость: {decimal_label(total)} тенге.', 13, 10, bold=True)
        elif provider == 'abdicul':
            pdf_field(page, (117, 435.4, 399, 453.7),
                      f'Изготовление стендов — {count} шт.', 11, 9)
            pdf_field(page, (117, 459, 550, 478),
                      f'Цена за 1 шт.: {decimal_label(stand_rate)} тенге', 11, 9)
            pdf_field(page, (405, 435.4, 550, 453.7),
                      f'{decimal_label(total)} тенге', 11, 7, align=fitz.TEXT_ALIGN_CENTER)
        else:
            raise ValueError(f'Нет обработчика подписанного КП: {mode}/{provider}')

        # Preserve searchable vector text and the original signature/stamp
        # pixels, not a screenshot of the entire page. Subset inserted fonts
        # to keep downloads small and page rendering fast.
        source.subset_fonts()
        source.save(destination, garbage=4, deflate=True, use_objstms=1)


def create_supplier_offer(mode: str, supplier: dict, destination: Path,
                          area: Decimal = Decimal(0), area_rate: Decimal = Decimal(0),
                          count: int = 0, stand_rate: Decimal = Decimal(0)) -> None:
    """All package types prefer their own signed PDF when one is registered."""
    provider = supplier['key']
    if supplier.get('signed_offer'):
        create_signed_offer(TEMPLATE_DIR / supplier['signed_offer'], destination,
                            mode, provider, area, area_rate, count, stand_rate)
        return
    source = TEMPLATE_DIR / supplier['offer']
    if source.suffix == '.pdf':
        if mode == 'stands':
            create_stands_offer_andrey(destination, count, stand_rate, source)
        elif mode == 'combined':
            create_combined_offer_archi(destination, area, count, stand_rate,
                                        money(area * area_rate) + money(count * stand_rate), area_rate, source)
        else:
            create_area_offer_archi(source, destination, area, area_rate)
    else:
        docx_path = destination.with_suffix('.docx')
        if mode == 'stands':
            updater = update_stands_offer_abdicul if provider == 'abdicul' else update_stands_offer_trade
            updater(docx_path, count, stand_rate, source)
        elif mode == 'combined':
            updater = update_combined_offer_best if provider == 'best' else update_combined_offer_nalibai
            updater(source, docx_path, area, count, stand_rate,
                    money(area * area_rate) + money(count * stand_rate), area_rate)
        else:
            update_area_offer(source, docx_path, provider, area, area_rate)
        convert_docx_to_pdf(docx_path, destination.parent)


def save_flattened_a4(source: fitz.Document, destination: Path, dpi: int = 144) -> None:
    """Save one edited supplier page as a simple, fast A4 image page."""
    page = source[0]
    a4 = fitz.paper_rect("a4")
    page_ratio = page.rect.width / page.rect.height
    a4_ratio = a4.width / a4.height
    if page_ratio >= a4_ratio:
        content_width = a4.width
        content_height = content_width / page_ratio
        left = 0
        top = (a4.height - content_height) / 2
    else:
        content_height = a4.height
        content_width = content_height * page_ratio
        left = (a4.width - content_width) / 2
        top = 0

    target_width = round(content_width * dpi / 72)
    render_scale = target_width / page.rect.width
    pixmap = page.get_pixmap(matrix=fitz.Matrix(render_scale, render_scale), alpha=False)
    flattened_page = pixmap.tobytes("jpeg", jpg_quality=90)

    normalized = fitz.open()
    out_page = normalized.new_page(width=a4.width, height=a4.height)
    out_page.insert_image(
        fitz.Rect(left, top, left + content_width, top + content_height),
        stream=flattened_page,
        keep_proportion=True,
    )
    normalized.save(
        destination,
        garbage=4,
        deflate=True,
        deflate_images=True,
        use_objstms=1,
    )
    normalized.close()








def add_labeled_line(paragraph, label: str, value: str, break_before: bool = False) -> None:
    if break_before:
        paragraph.add_run("\n")
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.name = "Times New Roman"
    value_run = paragraph.add_run(value)
    value_run.font.name = "Times New Roman"


def update_combined_specification(source: Path, destination: Path, area: Decimal, stand_count: int) -> None:
    doc = Document(source)
    for table in doc.tables:
        if len(table.rows) < 3 or len(table.columns) < 4:
            continue
        set_paragraph_text(table.cell(1, 3).paragraphs[0], decimal_label(area))
        set_paragraph_text(table.cell(2, 3).paragraphs[0], str(stand_count))
    doc.save(destination)


def update_combined_justification(source: Path, destination: Path, minimum_total: Decimal) -> None:
    doc = Document(source)
    value = decimal_label(minimum_total)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text
                    if "________" not in text:
                        continue
                    updated = re.sub(r"_{4,}", value, text)
                    if minimum_total > CAP_TOTAL:
                        updated = updated.replace("не превышает", "превышает")
                        updated = updated.replace("аспайды", "асады")
                    set_paragraph_text(paragraph, updated)
    doc.save(destination)


def update_combined_offer_best(
    source: Path,
    destination: Path,
    area: Decimal,
    stand_count: int,
    stand_rate: Decimal,
    total: Decimal,
    area_rate: Decimal,
) -> None:
    doc = Document(source)
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.startswith("Общая площадь:"):
            paragraph.clear()
            add_labeled_line(paragraph, "Общая площадь: ", f"{decimal_label(area)} м²")
            add_labeled_line(
                paragraph,
                "Стоимость работ за 1 м²: ",
                f"{decimal_label(area_rate)} тенге",
                break_before=True,
            )
            for run in paragraph.runs:
                run.font.size = Pt(11)
        elif paragraph.text.startswith("Стоимость стенда"):
            paragraph.clear()
            add_labeled_line(
                paragraph,
                "Стенды: ",
                f"{stand_count} шт. × {decimal_label(stand_rate)} = {decimal_label(stand_rate * stand_count)} тенге",
            )
            add_labeled_line(
                paragraph,
                "Общая стоимость: ",
                f"{decimal_label(total)} тенге",
                break_before=True,
            )
            for run in paragraph.runs:
                run.font.size = Pt(11)
        elif paragraph.text.strip().startswith("BEST PRODUCTION") or 19 <= paragraph_index <= 24:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(7.5)
            for run in paragraph.runs:
                run.font.size = Pt(6.8)
        elif 4 <= paragraph_index <= 16:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(11)
            for run in paragraph.runs:
                if run.font.size is None or run.font.size.pt > 10.5:
                    run.font.size = Pt(10.5)
    for paragraph_index, paragraph in reversed(list(enumerate(doc.paragraphs))):
        has_drawing = bool(paragraph._p.xpath(".//*[local-name()='drawing' or local-name()='pict']"))
        if paragraph_index > 2 and not paragraph.text.strip() and not has_drawing:
            paragraph._element.getparent().remove(paragraph._element)
    doc.save(destination)


def update_combined_offer_nalibai(
    source: Path,
    destination: Path,
    area: Decimal,
    stand_count: int,
    stand_rate: Decimal,
    total: Decimal,
    area_rate: Decimal,
) -> None:
    doc = Document(source)
    # The full-page background contains the gray sidebar. The stamp is a
    # separate inline image, so removing only page-sized background anchors
    # preserves the supplier's text, table, and stamp.
    section = doc.sections[0]
    for anchor in doc._element.xpath(".//wp:anchor[@behindDoc='1']"):
        extent = anchor.find(qn('wp:extent'))
        if extent is not None and (
            int(extent.get('cx', '0')) >= section.page_width * 0.9
            and int(extent.get('cy', '0')) >= section.page_height * 0.9
        ):
            drawing = anchor.getparent()
            drawing.getparent().remove(drawing)
    table = doc.tables[0]
    values = {
        2: (
            f"{decimal_label(area)} м² × {decimal_label(area_rate)} "
            f"тенге = {decimal_label(money(area_rate * area))} тенге"
        ),
        3: (
            f"{stand_count} шт. × {decimal_label(stand_rate)} тенге = "
            f"{decimal_label(stand_rate * stand_count)} тенге"
        ),
        4: f"{decimal_label(total)} тенге без НДС",
    }
    for row_index, value in values.items():
        paragraph = table.cell(row_index, 1).paragraphs[0]
        set_paragraph_text(paragraph, value)
        for run in paragraph.runs:
            run.font.size = Pt(10)
    for paragraph in reversed(doc.paragraphs):
        if paragraph.text.strip():
            break
        paragraph._element.getparent().remove(paragraph._element)
    doc.save(destination)


def create_combined_offer_archi(
    destination: Path,
    area: Decimal,
    stand_count: int,
    stand_rate: Decimal,
    total: Decimal,
    area_rate: Decimal,
    source_path: Path | None = None,
) -> None:
    source = fitz.open(source_path or TEMPLATE_DIR / 'combined/offer-archi.pdf')
    page = source[0]
    erase(
        page,
        fitz.Rect(400, 1892, 2180, 2034),
        fitz.Rect(400, 2060, 1950, 2150),
    )
    textbox(
        page,
        fitz.Rect(420, 1900, 2160, 1968),
        f"Отделочные работы: {decimal_label(area)} м² × {decimal_label(area_rate)} тенге",
        48.0,
        32.0,
    )
    textbox(
        page,
        fitz.Rect(420, 1960, 2210, 2035),
        (
            f"Стенды: {stand_count} шт. × {decimal_label(stand_rate)} = "
            f"{decimal_label(stand_rate * stand_count)} тенге"
        ),
        45.0,
        29.0,
    )
    textbox(
        page,
        fitz.Rect(420, 2070, 1900, 2142),
        f"Общая стоимость работ: {decimal_label(total)} тенге",
        47.0,
        30.0,
        bold=True,
    )
    save_flattened_a4(source, destination)
    source.close()


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not Path(soffice).exists():
        raise RuntimeError("LibreOffice не найден: без него нельзя собрать итоговый PDF")
    profile = output_dir / f".lo-profile-{uuid.uuid4().hex}"
    home = output_dir / f".lo-home-{uuid.uuid4().hex}"
    profile.mkdir()
    home.mkdir()
    env = os.environ.copy()
    env["HOME"] = str(home)
    command = [
        soffice,
        "--headless",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(docx_path),
    ]
    completed = subprocess.run(command, env=env, capture_output=True, text=True, timeout=90)
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.exists():
        detail = completed.stderr.strip() or completed.stdout.strip() or "неизвестная ошибка"
        raise RuntimeError(f"Не удалось преобразовать обоснование в PDF: {detail}")
    shutil.rmtree(profile, ignore_errors=True)
    shutil.rmtree(home, ignore_errors=True)
    return pdf_path


def merge_pdfs(destination: Path, sources: list[Path]) -> None:
    result = fitz.open()
    try:
        for source_path in sources:
            source = fitz.open(source_path)
            result.insert_pdf(source)
            source.close()
        result.set_metadata({"title": "Обоснование и коммерческие предложения", "author": "Генератор документов"})
        result.save(
            destination,
            garbage=4,
            deflate=True,
            deflate_images=True,
            deflate_fonts=True,
            use_objstms=1,
            compression_effort=50,
        )
    finally:
        result.close()


def update_stands_specification(destination: Path, count: int, source: Path | None = None) -> None:
    doc = Document(source or TEMPLATE_DIR / KTRU['stands']['specification'])
    for index, table in enumerate(doc.tables):
        quantity_cell = table.cell(1, 2)
        set_paragraph_text(quantity_cell.paragraphs[0], f"{count} {'штук' if index == 0 else 'дана'}")
        for paragraph in quantity_cell.paragraphs[1:]:
            paragraph._element.getparent().remove(paragraph._element)
        floating = table._tbl.tblPr.find(qn('w:tblpPr'))
        if floating is not None:
            table._tbl.tblPr.remove(floating)
        for row in table.rows:
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.height = None
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn('w:cantSplit')) is None:
                tr_pr.append(OxmlElement('w:cantSplit'))
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    # Remove source's blank spacer pages before the Kazakh part.
    for paragraph in list(doc.paragraphs):
        if not paragraph.text.strip() and not paragraph._p.xpath(".//*[local-name()='drawing' or local-name()='pict']"):
            paragraph._element.getparent().remove(paragraph._element)
        elif paragraph.text.strip() == 'ТЕХНИКАЛЫҚ СИПАТТАМА':
            paragraph.paragraph_format.page_break_before = True
        elif paragraph.text.strip() in ('Приложение 1', 'Қосымша 1'):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(3)
    doc.save(destination)


def update_stands_justification(destination: Path, count: int, rates: dict[str, Decimal], source: Path | None = None) -> None:
    doc = Document(source or TEMPLATE_DIR / KTRU['stands']['justification'])
    providers = (("Әбдіқұл", "abdicul"), ("Торговая сеть Казахстана", "trade"), ("Андрей и Ко", "andrey"))
    minimum = money(count * rates['abdicul'])
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text
        for name, key in providers:
            if f'«{name}»' in text and '_' in text:
                company = 'ТОО' if key == 'trade' else 'ИП'
                unit = 'шт.' if index < 23 else 'дана'
                set_paragraph_text(paragraph, f'{company} «{name}» — {decimal_label(rates[key])} тг × {count} {unit} = {decimal_label(money(rates[key] * count))} тг')
                break
        else:
            if '_' in text and ('По итогам' in text or 'Жүргізілген талдау' in text):
                updated = re.sub(r'_{4,}', decimal_label(minimum), text)
                if minimum > CAP_TOTAL:
                    updated = updated.replace('не превышает', 'превышает').replace('аспайды', 'асады')
                set_paragraph_text(paragraph, updated)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(12)
        if index in (0, 23):
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.size = Pt(14)
        if index == 23:
            paragraph.paragraph_format.page_break_before = True
    doc.save(destination)


def update_stands_offer_abdicul(destination: Path, count: int, rate: Decimal, source: Path | None = None) -> None:
    doc = Document(source or TEMPLATE_DIR / 'stands/offer-abdicul.docx')
    table = doc.tables[1]
    set_paragraph_text(table.cell(1, 1).paragraphs[0], f'Изготовление стендов — {count} шт.\n{decimal_label(rate)} тенге за 1 шт.')
    set_paragraph_text(table.cell(1, 2).paragraphs[0], f'{decimal_label(money(count * rate))} тенге')
    doc.save(destination)


def update_stands_offer_trade(destination: Path, count: int, rate: Decimal, source: Path | None = None) -> None:
    doc = Document(source or TEMPLATE_DIR / 'stands/offer-trade.docx')
    cell = doc.tables[1].cell(1, 1)
    set_paragraph_text(cell.paragraphs[0], f'{decimal_label(rate)} тг за 1 шт.\nКоличество: {count} шт.\n\nИтого: {decimal_label(money(count * rate))} тг\nБез НДС')
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    doc.save(destination)


def create_stands_offer_andrey(destination: Path, count: int, rate: Decimal, source_path: Path | None = None) -> None:
    source = fitz.open(source_path or TEMPLATE_DIR / 'stands/offer-andrey.pdf')
    page = source[0]
    erase(page, fitz.Rect(575, 2070, 1550, 2240), fitz.Rect(1580, 2070, 2235, 2240))
    textbox(page, fitz.Rect(585, 2080, 1545, 2230), f'Изготовление информационных стендов\nна тему «Жаңа Конституция» — {count} шт.', 44, 30)
    textbox(page, fitz.Rect(1600, 2080, 2220, 2230), f'{decimal_label(rate)} тг / шт.\nИтого: {decimal_label(money(count * rate))} тг', 44, 28, align=fitz.TEXT_ALIGN_CENTER)
    save_flattened_a4(source, destination)
    source.close()


def generate_stands_package(count: int, rates: dict[str, Decimal]) -> dict:
    kind = KTRU['stands']
    job_id = uuid.uuid4().hex[:12]
    job_dir = GENERATED_DIR / job_id
    job_dir.mkdir(parents=True)
    spec_name = f'Техническая_спецификация_стенды_{count}шт.docx'
    bundle_name = f'Обоснование_и_3_КП_стенды_{count}шт.pdf'
    update_stands_specification(job_dir / spec_name, count, TEMPLATE_DIR / kind['specification'])
    justification = job_dir / 'Обоснование.docx'
    update_stands_justification(justification, count, rates, TEMPLATE_DIR / kind['justification'])
    offer_pdfs = []
    for supplier in kind['suppliers']:
        pdf = job_dir / f'КП_{supplier["key"]}.pdf'
        create_supplier_offer('stands', supplier, pdf, count=count, stand_rate=rates[supplier['key']])
        offer_pdfs.append(pdf)
    merge_pdfs(job_dir / bundle_name, [convert_docx_to_pdf(justification, job_dir), *offer_pdfs])
    metadata = {
        'jobId': job_id,
        'ktru': KTRU['stands']['label'],
        'standCount': count,
        'files': [
            {'name': spec_name, 'kind': 'DOCX', 'description': 'Техническая спецификация на стенды', 'url': f'/download/{job_id}/{quote(spec_name)}'},
            {'name': bundle_name, 'kind': 'PDF', 'description': 'Обоснование и 3 коммерческих предложения', 'url': f'/download/{job_id}/{quote(bundle_name)}'},
        ],
    }
    (job_dir / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf-8')
    return metadata


def generate_combined_package(area: Decimal, stand_count: int, stand_rates: dict[str, Decimal], area_rates: dict[str, Decimal]) -> dict:
    kind = KTRU['combined']
    job_id = uuid.uuid4().hex[:12]
    job_dir = GENERATED_DIR / job_id
    job_dir.mkdir(parents=True)
    area_slug = slug_area(area)
    spec_name = f"Техническая_спецификация_{area_slug}м2_{stand_count}шт.docx"
    bundle_name = f"Обоснование_и_3_КП_{area_slug}м2_{stand_count}шт.pdf"
    spec_path = job_dir / spec_name
    bundle_path = job_dir / bundle_name

    justification_docx = job_dir / "Обоснование.docx"
    update_catalog_justification('combined', justification_docx, area, area_rates, stand_count, stand_rates)
    update_combined_specification(
        TEMPLATE_DIR / kind['specification'],
        spec_path,
        area,
        stand_count,
    )

    justification_pdf = convert_docx_to_pdf(justification_docx, job_dir)
    offer_pdfs = []
    for supplier in kind['suppliers']:
        provider = supplier['key']
        pdf = job_dir / f'КП_{provider}.pdf'
        create_supplier_offer('combined', supplier, pdf, area, area_rates[provider],
                              stand_count, stand_rates[provider])
        offer_pdfs.append(pdf)
    merge_pdfs(bundle_path, [justification_pdf, *offer_pdfs])

    metadata = {
        "jobId": job_id,
        "ktru": KTRU["combined"]["label"],
        "area": decimal_label(area),
        "standCount": stand_count,
        "files": [
            {
                "name": spec_name,
                "kind": "DOCX",
                "description": "Техспецификация с площадью и стендами",
                "url": f"/download/{job_id}/{quote(spec_name)}",
            },
            {
                "name": bundle_name,
                "kind": "PDF",
                "description": "Обоснование и 3 коммерческих предложения",
                "url": f"/download/{job_id}/{quote(bundle_name)}",
            },
        ],
    }
    (job_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return metadata


def generate_package(
    ktru_key: str,
    area: Decimal = Decimal(0),
    stand_count: int | None = None,
    stand_rates: dict[str, Decimal] | None = None,
    area_rates: dict[str, Decimal] | None = None,
) -> dict:
    if ktru_key not in KTRU:
        raise ValueError("Выберите вид закупки из списка")
    if ktru_key == 'stands':
        if stand_count is None or stand_rates is None:
            raise ValueError('Укажите количество и цену стенда')
        return generate_stands_package(stand_count, stand_rates)
    if ktru_key == "combined":
        if stand_count is None or stand_rates is None:
            raise ValueError("Укажите количество и цены стендов")
        defaults = supplier_rates(Decimal('25000'), (Decimal('7'), Decimal('15')), supplier_keys(ktru_key))
        return generate_combined_package(area, stand_count, stand_rates, area_rates or defaults)
    kind = KTRU[ktru_key]
    job_id = uuid.uuid4().hex[:12]
    job_dir = GENERATED_DIR / job_id
    job_dir.mkdir(parents=True)
    area_slug = slug_area(area)
    spec_name = f"Техническая_спецификация_{area_slug}м2.docx"
    bundle_name = f"Обоснование_и_3_КП_{area_slug}м2.pdf"
    spec_path = job_dir / spec_name
    bundle_path = job_dir / bundle_name

    area_rates = area_rates or supplier_rates(Decimal('25000'), (Decimal('7'), Decimal('15')), supplier_keys(ktru_key))
    justification_docx = job_dir / "Обоснование.docx"
    update_catalog_justification(ktru_key, justification_docx, area, area_rates)
    update_specification(TEMPLATE_DIR / kind['specification'], spec_path, area, kind)
    justification_pdf = convert_docx_to_pdf(justification_docx, job_dir)
    offer_pdfs = []
    for supplier in kind['suppliers']:
        provider = supplier['key']
        pdf_path = job_dir / f'КП_{provider}.pdf'
        create_supplier_offer(ktru_key, supplier, pdf_path, area, area_rates[provider])
        offer_pdfs.append(pdf_path)
    merge_pdfs(bundle_path, [justification_pdf, *offer_pdfs])

    metadata = {
        "jobId": job_id,
        "ktru": kind["label"],
        "area": decimal_label(area),
        "files": [
            {
                "name": spec_name,
                "kind": "DOCX",
                "description": "Редактируемая техническая спецификация",
                "url": f"/download/{job_id}/{quote(spec_name)}",
            },
            {
                "name": bundle_name,
                "kind": "PDF",
                "description": "Обоснование и 3 коммерческих предложения",
                "url": f"/download/{job_id}/{quote(bundle_name)}",
            },
        ],
    }
    (job_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return metadata


def generate_from_payload(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError('Некорректный запрос')
    ktru_key = str(payload.get('ktru', ''))
    if ktru_key not in KTRU:
        raise ValueError('Выберите вид закупки из списка')
    markups = (parse_markup(payload.get('markup1', '7')), parse_markup(payload.get('markup2', '15')))
    area = Decimal(0) if ktru_key == 'stands' else parse_area(payload.get('area', ''))
    keys = supplier_keys(ktru_key)
    area_rates = supplier_rates(parse_money(payload.get('areaPrice', '25000'), 'цену за м²'), markups, keys) if ktru_key != 'stands' else None
    count = None
    stand_rates = None
    if ktru_key in ('stands', 'combined'):
        count = parse_count(payload.get('standCount', ''))
        stand_rates = supplier_rates(parse_money(payload.get('standPrice', ''), 'цену за стенд'), markups, keys)
    return generate_package(ktru_key, area, count, stand_rates, area_rates)


class AppHandler(BaseHTTPRequestHandler):
    server_version = "DocumentGenerator/1.0"

    def log_message(self, format: str, *args) -> None:
        print(f"[{self.log_date_time_string()}] {format % args}")

    def send_bytes(self, body: bytes, content_type: str, status: int = 200, headers: dict | None = None) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        if headers:
            for key, value in headers.items():
                self.send_header(key, value)
        self.end_headers()
        self.wfile.write(body)

    def send_json(self, payload: dict, status: int = 200) -> None:
        self.send_bytes(json.dumps(payload, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8", status)

    def is_authenticated(self) -> bool:
        cookie = SimpleCookie()
        try:
            cookie.load(self.headers.get("Cookie", ""))
        except Exception:
            return False
        value = cookie.get(AUTH_COOKIE_NAME)
        return value is not None and hmac.compare_digest(value.value, AUTH_SESSION_TOKEN)

    def redirect(self, location: str, cookie: str | None = None) -> None:
        headers = {"Location": location, "Cache-Control": "no-store"}
        if cookie:
            headers["Set-Cookie"] = cookie
        self.send_bytes(b"", "text/plain; charset=utf-8", HTTPStatus.SEE_OTHER, headers)

    def serve_login(self, show_error: bool = False) -> None:
        path = STATIC_DIR / "login.html"
        if not path.exists():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        body = path.read_text(encoding="utf-8")
        if show_error:
            body = body.replace('id="login-error" class="login-error" hidden', 'id="login-error" class="login-error"')
        self.send_bytes(
            body.encode("utf-8"),
            "text/html; charset=utf-8",
            headers={"Cache-Control": "no-store"},
        )

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/health":
            self.send_json({"status": "ok"})
            return
        if parsed.path == "/login":
            if self.is_authenticated():
                self.redirect("/")
                return
            show_error = parse_qs(parsed.query).get("error") == ["1"]
            self.serve_login(show_error)
            return
        if parsed.path == "/":
            if not self.is_authenticated():
                self.redirect("/login")
                return
            self.send_bytes(
                render_index().encode('utf-8'),
                "text/html; charset=utf-8",
                headers={"Cache-Control": "no-store"},
            )
            return
        if parsed.path.startswith("/static/"):
            filename = parsed.path.removeprefix("/static/")
            if filename not in {"styles.css", "app.js"}:
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            content_types = {".css": "text/css; charset=utf-8", ".js": "application/javascript; charset=utf-8"}
            self.serve_file(STATIC_DIR / filename, content_types.get(Path(filename).suffix, "application/octet-stream"))
            return
        if parsed.path.startswith("/download/"):
            if not self.is_authenticated():
                self.redirect("/login")
                return
            parts = parsed.path.split("/", 3)
            if len(parts) != 4 or not re.fullmatch(r"[a-f0-9]{12}", parts[2]):
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            job_dir = GENERATED_DIR / parts[2]
            metadata_path = job_dir / "metadata.json"
            if not metadata_path.exists():
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            filename = unquote(parts[3])
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            allowed = {entry["name"] for entry in metadata["files"]}
            if filename not in allowed:
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            path = job_dir / filename
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if path.suffix == ".docx" else "application/pdf"
            disposition = f"attachment; filename*=UTF-8''{quote(filename)}"
            self.serve_file(path, mime, {"Content-Disposition": disposition})
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def serve_file(self, path: Path, content_type: str, headers: dict | None = None) -> None:
        if not path.exists() or not path.is_file():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        self.send_bytes(path.read_bytes(), content_type, headers=headers)

    def do_POST(self) -> None:
        path = urlparse(self.path).path
        if path == "/login":
            self.handle_login()
            return
        if path == "/logout":
            self.redirect(
                "/login",
                f"{AUTH_COOKIE_NAME}=; Path=/; HttpOnly; SameSite=Strict; Max-Age=0",
            )
            return
        if path != "/api/generate":
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        if not self.is_authenticated():
            self.send_json({"error": "Требуется авторизация"}, HTTPStatus.UNAUTHORIZED)
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 1_000_000:
                raise ValueError("Некорректный запрос")
            payload = json.loads(self.rfile.read(length))
            result = generate_from_payload(payload)
            self.send_json(result, HTTPStatus.CREATED)
        except (ValueError, json.JSONDecodeError) as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            self.send_json({"error": f"Не удалось собрать документы: {exc}"}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def handle_login(self) -> None:
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 10_000:
                raise ValueError
            fields = parse_qs(self.rfile.read(length).decode("utf-8"), keep_blank_values=True)
            username = fields.get("username", [""])[0]
            password = fields.get("password", [""])[0]
        except (ValueError, UnicodeDecodeError):
            self.redirect("/login?error=1")
            return

        valid_username = hmac.compare_digest(username.encode("utf-8"), AUTH_USERNAME.encode("utf-8"))
        valid_password = hmac.compare_digest(password.encode("utf-8"), AUTH_PASSWORD.encode("utf-8"))
        if not (valid_username and valid_password):
            self.redirect("/login?error=1")
            return

        cookie = (
            f"{AUTH_COOKIE_NAME}={AUTH_SESSION_TOKEN}; Path=/; HttpOnly; "
            f"SameSite=Strict; Max-Age={AUTH_SESSION_SECONDS}"
        )
        self.redirect("/", cookie)


def main() -> None:
    if not AUTH_USERNAME or not AUTH_PASSWORD:
        raise RuntimeError("Задайте переменные DOCGEN_USERNAME и DOCGEN_PASSWORD")
    railway_port = os.environ.get("PORT")
    host = os.environ.get("DOCGEN_HOST", "0.0.0.0" if railway_port else "127.0.0.1")
    port = int(railway_port or os.environ.get("DOCGEN_PORT", "8765"))
    server = ThreadingHTTPServer((host, port), AppHandler)
    print(f"Генератор документов запущен: http://{host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
